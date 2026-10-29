from io import BytesIO

from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .numbers import is_whole, unit_label
from .stock import StockReport

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r", "\n")
QTY_WHOLE_FORMAT = "#,##0"
QTY_FRACTION_FORMAT = "#,##0.0##"
MONEY_FORMAT = "#,##0.00"


def safe_text(value) -> str:
    """Neutralise CSV/Excel formula injection in user-supplied text."""
    text = str(value)
    return "'" + text if text.startswith(FORMULA_TRIGGERS) else text


def build_stock_workbook(report: StockReport) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _("Остатки")[:31]

    headers = [_("Товар"), _("Ед."), _("Остаток")]
    if report.with_prices:
        headers += [_("Цена"), _("Валюта"), _("Сумма")]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for row in report.rows:
        values = [safe_text(row["name"]), unit_label(row["unit"]), row["stock"]]
        if report.with_prices:
            values += [row["price"], row["currency"], row["amount"]]
        sheet.append(values)
        cells = sheet[sheet.max_row]
        cells[2].number_format = QTY_WHOLE_FORMAT if is_whole(row["stock"]) else QTY_FRACTION_FORMAT
        if report.with_prices:
            cells[3].number_format = MONEY_FORMAT
            cells[5].number_format = MONEY_FORMAT

    if report.with_prices and report.totals:
        sheet.append([])
        for code, amount in report.totals:
            sheet.append([_("Итого"), None, None, None, code, amount])
            cells = sheet[sheet.max_row]
            cells[5].number_format = MONEY_FORMAT
            for cell in cells:
                cell.font = Font(bold=True)

    widths = [48, 8, 14, 16, 10, 18]
    for index in range(len(headers)):
        sheet.column_dimensions[get_column_letter(index + 1)].width = widths[index]
    sheet.freeze_panes = "A2"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
