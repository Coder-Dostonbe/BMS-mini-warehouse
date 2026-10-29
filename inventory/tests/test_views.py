import io
import re
import zipfile
from decimal import Decimal
from pathlib import Path

from django.conf import settings
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from inventory import services
from inventory.models import Operation, Product

from .helpers import demo_user, key, make_product

SECRET_PRICE_FORMS = ["987654", "987 654", "987 654", "987,654"]


class PriceHidingTests(TestCase):
    def setUp(self):
        self.product = make_product(name="Секретный насос", unit="pcs", price="987654.32", currency="USD", stock="1")
        self.client.force_login(demo_user("storekeeper"))

    def assert_no_price(self, content: str):
        for needle in SECRET_PRICE_FORMS:
            self.assertNotIn(needle, content)

    def test_html_pages_have_no_price(self):
        for url in [
            reverse("inventory:stock"),
            reverse("inventory:stock") + "?archived=1",
            reverse("inventory:stock") + "?q=насос",
            reverse("inventory:receive"),
            reverse("inventory:issue"),
            reverse("inventory:receive") + f"?product={self.product.pk}",
        ]:
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, "Секретный насос")
                self.assert_no_price(response.content.decode())

    def test_context_has_no_price_either(self):
        response = self.client.get(reverse("inventory:stock"))
        row = response.context["report"].rows[0]
        self.assertNotIn("price", row)
        self.assertNotIn("amount", row)
        self.assertEqual(response.context["report"].totals, [])

    def test_error_pages_have_no_price(self):
        response = self.client.post(
            reverse("inventory:issue"),
            {"product": self.product.pk, "quantity": "5", "object_name": "Объект", "idempotency_key": key()},
        )
        self.assertContains(response, "На складе только 1 шт")
        self.assert_no_price(response.content.decode())

    def test_excel_has_no_price(self):
        response = self.client.get(reverse("inventory:export"))
        self.assertEqual(response.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            for name in archive.namelist():
                self.assert_no_price(archive.read(name).decode("utf-8", errors="replace"))
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual([c.value for c in sheet[1]], ["Товар", "Ед.", "Остаток"])

    def test_director_does_see_it(self):
        self.client.force_login(demo_user("director"))
        self.assertContains(self.client.get(reverse("inventory:stock")), "987 654,32 $")


class FormInputTests(TestCase):
    def setUp(self):
        self.director = demo_user("director")
        self.client.force_login(self.director)
        self.cable = make_product(name="Кабель", unit="m", price="1.00")
        self.bolts = make_product(name="Болт", unit="pcs", price="1.00")

    def receive(self, product, quantity, **extra):
        data = {"product": product.pk, "quantity": quantity, "idempotency_key": key(), **extra}
        return self.client.post(reverse("inventory:receive"), data)

    def test_comma_and_spaces_work(self):
        self.assertEqual(self.receive(self.cable, "1 000,5").status_code, 302)
        self.cable.refresh_from_db()
        self.assertEqual(self.cable.stock, Decimal("1000.5"))

    def test_fractional_pieces_rejected(self):
        response = self.receive(self.bolts, "1,5")
        self.assertContains(response, "Для штук количество должно быть целым")
        self.assertContains(response, 'value="1,5"')
        self.assertEqual(Operation.objects.count(), 0)

    def test_bad_numbers_are_form_errors_not_crashes(self):
        for raw in ["", "NaN", "Infinity", "1e5", "-1", "0", "abc", "1,2345", "99999999999999"]:
            with self.subTest(raw=raw):
                response = self.receive(self.cable, raw)
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, "errorlist")
        self.assertEqual(Operation.objects.count(), 0)

    def test_quantity_input_mode_hints(self):
        response = self.client.get(reverse("inventory:receive"))
        self.assertContains(response, 'inputmode="decimal"')
        self.assertContains(response, 'data-unit="pcs"')
        self.assertContains(response, "Болт — 0 шт")

    def test_same_key_posted_twice_creates_one_operation(self):
        data = {"product": self.cable.pk, "quantity": "7", "idempotency_key": key()}
        first = self.client.post(reverse("inventory:receive"), data, follow=True)
        second = self.client.post(reverse("inventory:receive"), data, follow=True)
        self.assertContains(first, "Приход сохранён")
        self.assertContains(second, "Операция уже сохранена")
        self.assertEqual(Operation.objects.count(), 1)
        self.cable.refresh_from_db()
        self.assertEqual(self.cable.stock, Decimal("7"))

    def test_form_key_survives_validation_error(self):
        k = key()
        response = self.client.post(reverse("inventory:receive"), {"product": self.cable.pk, "quantity": "x", "idempotency_key": k})
        self.assertContains(response, f'value="{k}"')

    def test_missing_or_broken_key_is_a_stale_form_error(self):
        for bad in [None, "", "not-a-uuid"]:
            data = {"product": self.cable.pk, "quantity": "1"}
            if bad is not None:
                data["idempotency_key"] = bad
            with self.subTest(key=bad):
                response = self.client.post(reverse("inventory:receive"), data)
                self.assertContains(response, "Форма устарела")
                self.assertRegex(response.content.decode(), r'name="idempotency_key" value="[0-9a-f-]{36}"')
        self.assertEqual(Operation.objects.count(), 0)

    def test_issue_without_object_is_rejected(self):
        services.receive(self.director, self.cable, Decimal("5"), key=key())
        response = self.client.post(
            reverse("inventory:issue"), {"product": self.cable.pk, "quantity": "1", "object_name": "   ", "idempotency_key": key()}
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Operation.objects.filter(type="out").count(), 0)

    def test_page_reload_after_save_is_a_get(self):
        response = self.client.post(reverse("inventory:receive"), {"product": self.cable.pk, "quantity": "1", "idempotency_key": key()})
        self.assertRedirects(response, reverse("inventory:stock"))

    def test_preselected_product_and_once_guard(self):
        response = self.client.get(reverse("inventory:receive") + f"?product={self.cable.pk}")
        self.assertRegex(response.content.decode(), rf'<option value="{self.cable.pk}"[^>]*selected')
        self.assertContains(response, "data-once")


class No500Tests(TestCase):
    def setUp(self):
        self.client.force_login(demo_user("director"))
        self.product = make_product(name="Кабель", stock="5")
        self.archived = make_product(name="Старьё", unit="pcs", price="1.00")
        services.archive_product(demo_user("director"), self.archived)

    def test_garbage_requests(self):
        huge = "9" * 30
        cases = [
            ("get", f"/products/{huge}/edit/", None, 404),
            ("get", "/products/999999/edit/", None, 404),
            ("post", "/products/999999/archive/", {}, 404),
            ("get", "/history/999999/cancel/", None, 404),
            ("get", f"/history/{huge}/cancel/", None, 404),
            ("get", "/no-such-page/", None, 404),
            ("get", "/history/?product=abc&type=zzz&page=abc", None, 200),
            ("get", f"/history/?product={huge}&page={huge}", None, 200),
            ("get", "/history/?page=-1", None, 200),
            ("get", "/?q=" + "я" * 5000, None, 200),
            ("get", f"/receive/?product={huge}", None, 200),
            ("post", "/receive/", {"product": "abc", "quantity": "1", "idempotency_key": key()}, 200),
            ("post", "/receive/", {"product": "999999", "quantity": "1", "idempotency_key": key()}, 200),
            ("post", "/receive/", {"product": huge, "quantity": "1", "idempotency_key": key()}, 200),
            ("post", "/receive/", {"product": self.archived.pk, "quantity": "1", "idempotency_key": key()}, 200),
            ("post", "/issue/", {"product": self.product.pk, "quantity": "1e3", "object_name": "x", "idempotency_key": key()}, 200),
            ("post", "/issue/", {}, 200),
            ("post", "/products/new/", {"name": "X", "unit": "kg", "price": "1e5", "currency": "EUR"}, 200),
            ("post", "/products/new/", {}, 200),
            ("post", f"/products/{self.product.pk}/edit/", {"name": "", "unit": "m", "price": "-1", "currency": "USD"}, 200),
        ]
        for method, url, data, status in cases:
            with self.subTest(method=method, url=url[:80]):
                response = getattr(self.client, method)(url, data or {})
                self.assertEqual(response.status_code, status)
        self.assertEqual(Operation.objects.count(), 1)

    def test_custom_404_page(self):
        response = self.client.get("/no-such-page/")
        self.assertContains(response, "Страница не найдена", status_code=404)
        self.assertContains(response, "На главную", status_code=404)


class ProductPageTests(TestCase):
    def setUp(self):
        self.director = demo_user("director")
        self.client.force_login(self.director)

    def test_unit_cannot_be_changed_through_the_form_after_movements(self):
        product = make_product(name="Кабель", unit="m", stock="3")
        response = self.client.get(reverse("inventory:product_edit", args=[product.pk]))
        self.assertContains(response, "Нельзя изменить: по товару уже есть операции")
        response = self.client.post(
            reverse("inventory:product_edit", args=[product.pk]),
            {"name": "Кабель", "unit": "pcs", "price": "1", "currency": "USD"},
        )
        self.assertEqual(response.status_code, 302)
        product.refresh_from_db()
        self.assertEqual(product.unit, "m")

    def test_product_with_movements_cannot_be_deleted_through_post(self):
        product = make_product(name="Кабель", stock="3")
        response = self.client.post(reverse("inventory:product_delete", args=[product.pk]), follow=True)
        self.assertContains(response, "Нельзя удалить товар")
        self.assertTrue(Product.objects.filter(pk=product.pk).exists())

    def test_duplicate_name_is_a_form_error(self):
        make_product(name="Кабель")
        response = self.client.post(
            reverse("inventory:product_create"), {"name": " КАБЕЛЬ ", "unit": "m", "price": "1", "currency": "USD"}
        )
        self.assertContains(response, "Товар с таким названием уже есть")
        self.assertEqual(Product.objects.count(), 1)

    def test_stock_is_not_editable_through_the_form(self):
        product = make_product(name="Кабель", stock="3")
        self.client.post(
            reverse("inventory:product_edit", args=[product.pk]),
            {"name": "Кабель", "unit": "m", "price": "1", "currency": "USD", "stock": "999"},
        )
        product.refresh_from_db()
        self.assertEqual(product.stock, Decimal("3"))

    def test_archive_and_restore(self):
        product = make_product(name="Лента")
        self.client.post(reverse("inventory:product_archive", args=[product.pk]))
        product.refresh_from_db()
        self.assertTrue(product.is_archived)
        response = self.client.get(reverse("inventory:receive"))
        self.assertNotContains(response, "Лента — ")
        response = self.client.get(reverse("inventory:stock"))
        self.assertNotContains(response, "<td data-label=\"Товар\">Лента</td>", html=False)
        response = self.client.get(reverse("inventory:stock") + "?archived=1")
        self.assertContains(response, "Лента")
        self.client.post(reverse("inventory:product_restore", args=[product.pk]))
        product.refresh_from_db()
        self.assertFalse(product.is_archived)

    def test_archive_filter_is_director_only(self):
        product = make_product(name="Лента")
        services.archive_product(self.director, product)
        self.client.force_login(demo_user("manager"))
        response = self.client.get(reverse("inventory:stock") + "?archived=1")
        self.assertNotContains(response, "Лента")


class SearchTests(TestCase):
    def setUp(self):
        self.client.force_login(demo_user("manager"))
        make_product(name="Кабель ВВГ")
        make_product(name="КАБЕЛЬ медный")
        make_product(name="Cable UTP")
        make_product(name="Ёлочный выключатель")

    def names(self, query):
        response = self.client.get(reverse("inventory:stock"), {"q": query})
        return sorted(row["name"] for row in response.context["report"].rows)

    def test_cyrillic_search_ignores_case(self):
        self.assertEqual(self.names("кабель"), sorted(["Кабель ВВГ", "КАБЕЛЬ медный"]))
        self.assertEqual(self.names("КАБ"), sorted(["Кабель ВВГ", "КАБЕЛЬ медный"]))
        self.assertEqual(self.names("МЕДН"), ["КАБЕЛЬ медный"])
        self.assertEqual(self.names("ёлоч"), ["Ёлочный выключатель"])
        self.assertEqual(self.names("cable"), ["Cable UTP"])
        self.assertEqual(self.names("нет такого"), [])
        self.assertEqual(len(self.names("")), 4)

    def test_search_wildcards_are_literal(self):
        self.assertEqual(self.names("%"), [])
        self.assertEqual(self.names("_"), [])


class HistoryTests(TestCase):
    def setUp(self):
        self.director = demo_user("director")
        self.client.force_login(self.director)
        self.cable = make_product(name="Кабель")
        self.bolt = make_product(name="Болт", unit="pcs")

    def test_pagination_by_fifty(self):
        for _ in range(55):
            services.receive(self.director, self.bolt, Decimal("1"), key=key())
        page1 = self.client.get(reverse("inventory:history"))
        page2 = self.client.get(reverse("inventory:history"), {"page": 2})
        self.assertEqual(len(page1.context["events"]), 50)
        self.assertEqual(len(page2.context["events"]), 5)
        self.assertContains(page1, "Страница 1 из 2")

    def test_filters(self):
        services.receive(self.director, self.cable, Decimal("5"), key=key())
        services.receive(self.director, self.bolt, Decimal("5"), key=key())
        out = services.issue(self.director, self.bolt, Decimal("1"), "Объект", key=key()).operation
        services.cancel(self.director, out)

        def events(**params):
            response = self.client.get(reverse("inventory:history"), params)
            return [(e["kind"], e["op"].product.name, e["op"].type) for e in response.context["events"]]

        self.assertEqual(len(events()), 4)
        self.assertEqual(
            events(product=self.bolt.pk),
            [("cancel", "Болт", "out"), ("operation", "Болт", "out"), ("operation", "Болт", "in")],
        )
        self.assertEqual(events(type="in"), [("operation", "Болт", "in"), ("operation", "Кабель", "in")])
        self.assertEqual(events(product=self.cable.pk, type="out"), [])

    def test_cancel_button_only_for_active_operations_and_director(self):
        op = services.receive(self.director, self.cable, Decimal("5"), key=key()).operation
        cancel_url = reverse("inventory:cancel", args=[op.pk])
        self.assertContains(self.client.get(reverse("inventory:history")), cancel_url)
        self.client.force_login(demo_user("manager"))
        self.assertNotContains(self.client.get(reverse("inventory:history")), cancel_url)
        self.client.force_login(self.director)
        services.cancel(self.director, op)
        self.assertNotContains(self.client.get(reverse("inventory:history")), cancel_url)

    def test_cancel_confirmation_page_changes_nothing(self):
        op = services.receive(self.director, self.cable, Decimal("5"), key=key()).operation
        response = self.client.get(reverse("inventory:cancel", args=[op.pk]))
        self.assertContains(response, "Отменить операцию")
        self.assertContains(response, "data-confirm")
        op.refresh_from_db()
        self.assertFalse(op.is_cancelled)

    def test_second_cancel_through_post_is_refused(self):
        op = services.receive(self.director, self.cable, Decimal("5"), key=key()).operation
        self.client.post(reverse("inventory:cancel", args=[op.pk]))
        response = self.client.post(reverse("inventory:cancel", args=[op.pk]))
        self.assertContains(response, "Операция уже отменена (director,")
        self.cable.refresh_from_db()
        self.assertEqual(self.cable.stock, Decimal("0"))

    def test_history_cannot_be_deleted(self):
        op = services.receive(self.director, self.cable, Decimal("5"), key=key()).operation
        self.assertEqual(self.client.delete(reverse("inventory:history")).status_code, 405)
        self.assertEqual(self.client.delete(reverse("inventory:cancel", args=[op.pk])).status_code, 405)
        self.assertEqual(self.client.post(f"/history/{op.pk}/").status_code, 404)
        self.assertEqual(self.client.post(f"/history/{op.pk}/delete/").status_code, 404)
        self.assertEqual(Operation.objects.count(), 1)


class ExcelTests(TestCase):
    def setUp(self):
        self.director = demo_user("director")
        self.client.force_login(self.director)

    def workbook(self, **params):
        response = self.client.get(reverse("inventory:export"), params)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertRegex(response["Content-Disposition"], r'attachment; filename="stock_\d{4}-\d{2}-\d{2}\.xlsx"')
        return load_workbook(io.BytesIO(response.content)).active

    def test_numbers_are_numeric_cells_with_formats_and_totals(self):
        make_product(name="Кабель", unit="m", price="0.80", currency="USD", stock="69.5")
        make_product(name="Болт", unit="pcs", price="0.10", currency="USD", stock="3")
        make_product(name="Труба", unit="m", price="12500", currency="UZS", stock="1000")
        sheet = self.workbook()
        self.assertEqual([c.value for c in sheet[1]], ["Товар", "Ед.", "Остаток", "Цена", "Валюта", "Сумма"])
        rows = {row[0].value: row for row in sheet.iter_rows(min_row=2) if row[0].value}
        bolt = rows["Болт"]
        self.assertEqual(bolt[2].data_type, "n")
        self.assertEqual(bolt[2].value, 3)
        self.assertEqual(bolt[3].number_format, "#,##0.00")
        self.assertAlmostEqual(bolt[5].value, 0.30)
        self.assertAlmostEqual(rows["Кабель"][2].value, 69.5)
        self.assertEqual(rows["Кабель"][2].number_format, "#,##0.0##")
        self.assertAlmostEqual(rows["Кабель"][5].value, 55.60)
        totals = [(row[4].value, row[5].value) for row in sheet.iter_rows(min_row=2) if row[0].value == "Итого"]
        self.assertEqual(len(totals), 2)
        self.assertEqual(totals[0][0], "UZS")
        self.assertAlmostEqual(totals[0][1], 12500000)
        self.assertEqual(totals[1][0], "USD")
        self.assertAlmostEqual(totals[1][1], 55.90)

    def test_formula_injection_is_neutralised(self):
        evil = ["=HYPERLINK(\"http://x\")", "+1+1", "-2+3", "@SUM(A1)", "=cmd|'/C calc'!A0"]
        for name in evil:
            make_product(name=name)
        sheet = self.workbook()
        values = [row[0] for row in sheet.iter_rows(min_row=2) if row[0].value and row[0].value != "Итого"]
        self.assertEqual(len(values), len(evil))
        for cell in values:
            with self.subTest(value=cell.value):
                self.assertEqual(cell.data_type, "s")
                self.assertTrue(cell.value.startswith("'"))

    def test_search_filter_is_applied(self):
        make_product(name="Кабель")
        make_product(name="Болт", unit="pcs")
        sheet = self.workbook(q="кабель")
        names = [row[0].value for row in sheet.iter_rows(min_row=2) if row[0].value and row[0].value != "Итого"]
        self.assertEqual(names, ["Кабель"])


class StaticSafetyTests(TestCase):
    def test_no_safe_filter_or_mark_safe_anywhere(self):
        root = Path(settings.BASE_DIR)
        for path in [*root.glob("templates/**/*.html"), *root.glob("inventory/**/*.py"), *root.glob("accounts/**/*.py")]:
            if "tests" in path.parts:
                continue
            text = path.read_text(encoding="utf-8")
            with self.subTest(path=str(path.relative_to(root))):
                self.assertIsNone(re.search(r"\|\s*safe\b|mark_safe|autoescape\s+off", text))

    def test_xss_in_product_name_is_escaped(self):
        make_product(name="<script>alert(1)</script>")
        self.client.force_login(demo_user("director"))
        response = self.client.get(reverse("inventory:stock"))
        self.assertNotContains(response, "<script>alert(1)</script>")
        self.assertContains(response, "&lt;script&gt;alert(1)&lt;/script&gt;")
